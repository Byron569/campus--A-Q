"""生成模拟校园资料。

设计依据：docs/03-开发任务清单.md §1.6（M1-16）

为什么需要：真实校园资料尚未到位（docs/01 §10.4「数据：待补」），
先用构造的资料开发与验证入库、检索链路。

产出：`data/samples/<分类>/` 共 8 份，格式覆盖 `.md` / `.txt` / `.docx`，
按分类分子目录，便于用 `ingest_cli.py` 分批带上正确的 `--category` 入库。

说明：不生成 PDF。设计选型里 pypdf 只用于读取，写 PDF 需要再引入第三方库，
为样例数据加依赖不划算；PDF 解析分支已由 tests/test_loader.py 覆盖。

用法：
    python scripts/make_samples.py
"""

from __future__ import annotations

import logging
import sys
from pathlib import Path

# 直接以 `python scripts/xxx.py` 运行时，sys.path[0] 是 scripts/ 而非项目根，
# 因此需要手动把项目根加进来才能 import config / src
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from config.settings import get_settings  # noqa: E402

logger = logging.getLogger("scripts.make_samples")

# 内容一律写成「可检索的具体事实」（时间、地点、金额、比例、流程），
# 便于验证检索是否真的召回正确段落，而不是只看是否返回结果。
FRESHMAN_REPORT = """# 新生报到须知

## 一、报到时间与地点

报到时间为 2026 年 9 月 1 日至 9 月 2 日，每天 8:30 至 17:30。
报到地点为学校体育馆（南门进入后直行 200 米）。

## 二、需携带材料

1. 录取通知书原件
2. 本人身份证原件及复印件 2 份
3. 一寸免冠照片 4 张

## 三、缴费方式

学费与住宿费统一通过「校园统一支付平台」在线缴纳，
不支持现场现金缴费。缴费截止日期为 2026 年 9 月 10 日。

## 四、特殊情况

因故不能按期报到的，须提前向所在学院提交书面请假申请，
请假期限原则上不超过两周。未经请假逾期两周未报到的，
按自动放弃入学资格处理。
"""

DORM_MOVE_PARAGRAPHS = [
    ("h1", "宿舍搬迁通知"),
    ("p", "为配合宿舍楼检修，学校将于 2026 年 9 月 20 日至 9 月 22 日组织集中搬迁。"),
    ("h2", "一、搬迁申请"),
    (
        "p",
        "需要搬迁宿舍的同学，必须提前三个工作日向辅导员提交书面申请，"
        "经学生工作处审批同意后方可搬迁。未提交申请不得自行调换床位。",
    ),
    ("h2", "二、搬迁流程"),
    (
        "p",
        "第一步，向辅导员领取《宿舍搬迁申请单》；第二步，原宿舍楼管理员核验物品；"
        "第三步，到新宿舍楼管理员处登记并领取钥匙。",
    ),
    ("h2", "三、办理时间与地点"),
    (
        "table",
        [
            ["事项", "时间", "地点"],
            ["领取申请单", "工作日 9:00 至 16:00", "各学院辅导员办公室"],
            ["物品核验", "搬迁当日 8:00 至 11:00", "原宿舍楼一层值班室"],
            ["登记领钥匙", "搬迁当日 14:00 至 17:00", "新宿舍楼一层值班室"],
        ],
    ),
    ("h2", "四、注意事项"),
    (
        "p",
        "搬迁期间个人物品由本人负责，学校不承担保管责任。"
        "空调、热水器等公共设施不得私自拆卸带走。",
    ),
]

DORM_RULES = """第一章 总则

第一条 为规范学生宿舍管理，营造安全、整洁的住宿环境，制定本规定。

第二章 作息与门禁

第二条 宿舍楼门禁时间为每日 23:00，超过门禁时间返回的须在值班室登记晚归，并说明原因。

第三条 未经批准不得留宿他人。因亲友探访需临时住宿的，应提前一天向辅导员报备。

第三章 用电安全

第四条 宿舍内禁止使用功率超过 800 瓦的电器，包括电热毯、电炉、热得快等。

第五条 严禁私拉乱接电线。因违规用电造成事故的，除赔偿损失外，还将给予纪律处分。

第四章 卫生检查

第六条 学生工作处每周三下午组织卫生检查，检查结果计入宿舍综合评定。
"""

CAMPUS_CARD = """# 校园卡使用说明

## 一、功能范围

校园卡可用于食堂消费、图书馆借阅、宿舍门禁与机房上机，
是校内唯一通用的身份与消费凭证。

## 二、挂失与解挂

校园卡遗失后，应在 24 小时内通过「校园一卡通」自助终端或线上服务大厅办理挂失。
挂失满 24 小时后仍找回的，可持本人身份证到服务大厅办理解挂。

## 三、补办

补办地点为行政楼一楼服务大厅，办理时间为工作日 8:30 至 17:00。
补办需携带本人身份证，工本费 20 元，一般当场可取。

## 四、初始密码

校园卡初始消费密码为本人身份证号码后六位，首次使用后请及时修改。
"""

LIBRARY_RULES = """# 图书馆借阅规则

## 一、借阅权限

本科生每人可同时借阅图书 10 册，研究生可借阅 20 册，教师可借阅 30 册。

## 二、借阅期限

普通图书借期 30 天，可续借 2 次，每次续借 15 天。
续借可通过图书馆微信公众号或自助借还机办理，已被他人预约的图书不可续借。

## 三、逾期与赔偿

逾期归还的，每册每天收取 0.1 元滞纳金，滞纳金累计上限为该书原价的 2 倍。
图书遗失的，应购买同名同版本图书赔偿，并缴纳 5 元加工费。

## 四、开放时间

借阅区开放时间为周一至周日 8:00 至 22:00，法定节假日另行通知。
自习区全天开放。
"""

LEAVE_PARAGRAPHS = [
    ("h1", "学生请假与考勤管理办法"),
    ("h2", "一、请假类型"),
    (
        "p",
        "请假分为病假、事假与公假三类。病假须附校医院或二级以上医院出具的诊断证明；"
        "事假须说明具体事由；公假指代表学校参加竞赛、会议等活动。",
    ),
    ("h2", "二、审批权限"),
    (
        "table",
        [
            ["请假天数", "审批人", "需提交材料"],
            ["1 天以内", "辅导员", "请假条"],
            ["2 至 3 天", "辅导员核实，学院备案", "请假条 + 相关证明"],
            ["超过 3 天", "学院分管领导审批", "请假条 + 相关证明 + 家长确认"],
        ],
    ),
    ("h2", "三、销假"),
    ("p", "请假期满后应于 1 个工作日内到辅导员处销假，逾期未销假的按旷课处理。"),
    ("h2", "四、考勤处理"),
    (
        "p",
        "未经批准缺勤的按旷课处理。一学期内累计旷课达 10 学时的，给予警告处分；"
        "达 20 学时的，给予严重警告处分；情节更重的按学校学籍管理规定处理。",
    ),
]

SCHOLARSHIP = """第一章 总则

第一条 为激励学生勤奋学习、全面发展，根据国家有关规定，结合本校实际，制定本细则。

第二章 奖项与标准

第二条 国家奖学金标准为每人每学年 8000 元，奖励名额由国家下达。

第三条 国家励志奖学金标准为每人每学年 5000 元，面向家庭经济困难且品学兼优的学生。

第四条 校级学业奖学金分三等：一等 3000 元，二等 2000 元，三等 1000 元。

第三章 评定办法

第五条 评定时间为每年 9 月，评定周期为上一学年。

第六条 学业成绩在综合测评中的权重不低于 70%，其余为思想品德与社会实践表现。

第七条 同一年度内，国家奖学金与国家励志奖学金不可兼得。
"""

DATA_STRUCTURE_HOMEWORK = """# 数据结构课程作业说明

## 一、课程信息

课程名称：数据结构与算法；授课教师：计算机学院 陈老师。
教材为《数据结构（C 语言版）》第 2 版。

## 二、作业安排

本学期共 6 次编程作业，覆盖线性表、栈与队列、树、图、排序与查找。
每次作业需提交源代码与一份实验报告。

## 三、提交要求

提交截止时间为每周日 22:00，通过课程平台提交，不接受邮件提交。
迟交的每次扣减本次成绩的 20%，迟交超过 3 天的不予批改。

## 四、评分与抄袭

作业成绩占课程总评的 40%，期末考试占 60%。
代码查重率超过 30% 的，本次作业记零分，情节严重者按学校学术规范处理。
"""


def _write_text(target: Path, content: str) -> None:
    target.write_text(content, encoding="utf-8")


def _write_docx(target: Path, blocks: list[tuple]) -> None:
    """按 (类型, 内容) 序列写 DOCX：h1/h2 为标题，p 为段落，table 为表格。

    表格是刻意保留的：校园通知里时间表、审批权限表大量以表格承载，
    只取段落会静默丢内容（见 loader 的 DOCX 抽取说明）。
    """
    import docx

    document = docx.Document()
    for kind, payload in blocks:
        if kind == "h1":
            document.add_heading(payload, level=1)
        elif kind == "h2":
            document.add_heading(payload, level=2)
        elif kind == "p":
            document.add_paragraph(payload)
        elif kind == "table":
            rows = payload
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                for col_index, cell in enumerate(row):
                    table.cell(row_index, col_index).text = cell
        else:  # pragma: no cover - 内置数据写错时的兜底
            raise ValueError(f"未知的 DOCX 块类型：{kind}")
    document.save(str(target))


def build_samples() -> list[tuple[str, str, str, object]]:
    """返回 (分类, 文件名, 格式, 内容) 列表。"""
    return [
        ("freshman", "新生报到须知.md", "md", FRESHMAN_REPORT),
        ("freshman", "宿舍搬迁通知.docx", "docx", DORM_MOVE_PARAGRAPHS),
        ("freshman", "宿舍管理规定.txt", "txt", DORM_RULES),
        ("freshman", "校园卡使用说明.md", "md", CAMPUS_CARD),
        ("admin", "图书馆借阅规则.md", "md", LIBRARY_RULES),
        ("admin", "请假与考勤管理办法.docx", "docx", LEAVE_PARAGRAPHS),
        ("admin", "奖学金评定细则.txt", "txt", SCHOLARSHIP),
        ("course", "数据结构课程作业说明.md", "md", DATA_STRUCTURE_HOMEWORK),
    ]


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")

    settings = get_settings()
    settings.ensure_dirs()
    root = settings.samples_path

    samples = build_samples()
    for category, filename, kind, content in samples:
        target = root / category / filename
        target.parent.mkdir(parents=True, exist_ok=True)
        if kind == "docx":
            _write_docx(target, content)  # type: ignore[arg-type]
        else:
            _write_text(target, content)  # type: ignore[arg-type]
        logger.info("已生成 %s（%s，%d 字节）", target.relative_to(root.parent.parent), category, target.stat().st_size)

    print(f"\n共生成 {len(samples)} 份模拟资料，目录：{root}")
    print("按分类入库（公共文档）：")
    for category in ("freshman", "admin", "course"):
        print(f"  python scripts/ingest_cli.py {root / category} --public --category {category}")
    print("检索验证：")
    print('  python scripts/ingest_cli.py --query "搬宿舍需要提前申请吗"')
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
