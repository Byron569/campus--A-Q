"""文档解析单元测试。

覆盖 docs/03 §6.1 相关用例，以及 docs/02 §11 异常处理策略表中与解析相关的分支：
文件类型不支持 / 文件为空 / 图片型 PDF / 加密 PDF / 文件不存在。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pypdf import PdfWriter

from src.errors import ParseError, UnsupportedFileType
from src.ingest.loader import clean_text, load_document


# ---------------- 辅助：构造测试文件 ----------------


def make_txt(tmp_path: Path, content: str, name: str = "notice.txt", encoding: str = "utf-8") -> Path:
    target = tmp_path / name
    target.write_text(content, encoding=encoding)
    return target


def make_blank_pdf(tmp_path: Path, pages: int = 1, name: str = "scan.pdf") -> Path:
    """生成无文本的 PDF —— 模拟扫描件。"""
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    target = tmp_path / name
    with open(target, "wb") as fh:
        writer.write(fh)
    return target


def make_encrypted_pdf(tmp_path: Path, name: str = "locked.pdf") -> Path:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("secret")
    target = tmp_path / name
    with open(target, "wb") as fh:
        writer.write(fh)
    return target


def make_docx(tmp_path: Path, name: str = "handbook.docx") -> Path:
    import docx

    document = docx.Document()
    document.add_paragraph("第二章 宿舍搬迁")
    document.add_paragraph("学生因转专业等原因需搬迁的，应提前三个工作日提交申请。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "报到时间"
    table.cell(0, 1).text = "九月一日至九月二日"
    target = tmp_path / name
    document.save(str(target))
    return target


# ---------------- clean_text ----------------


def test_clean_text_trims_and_collapses_blank_lines() -> None:
    raw = "  第一章   \n\n\n\n   第一条 规定。   \n   \n"
    assert clean_text(raw) == "第一章\n\n第一条 规定。"


def test_clean_text_on_empty_input() -> None:
    assert clean_text("") == ""
    assert clean_text("   \n\n  ") == ""


# ---------------- TXT / MD ----------------


def test_load_txt(tmp_path: Path) -> None:
    path = make_txt(tmp_path, "第一章 总则\n\n第一条 为规范宿舍管理，制定本办法。")
    result = load_document(path)

    assert not result.is_empty
    assert result.needs_ocr is False
    assert result.warning == ""
    assert len(result.sections) == 1
    assert result.sections[0].page is None
    assert "为规范宿舍管理" in result.sections[0].text


def test_load_markdown(tmp_path: Path) -> None:
    path = make_txt(tmp_path, "# 新生报到须知\n\n报到时间：九月一日。", name="guide.md")
    result = load_document(path)

    assert not result.is_empty
    assert "新生报到须知" in result.sections[0].text


def test_load_gbk_encoded_txt(tmp_path: Path) -> None:
    """中文校园文档常见 GBK 编码，应能兜底读取。"""
    path = make_txt(tmp_path, "校园卡遗失后应及时挂失。", name="gbk.txt", encoding="gbk")
    result = load_document(path)

    assert "校园卡遗失后应及时挂失" in result.sections[0].text


# ---------------- 空文件 ----------------


def test_empty_txt_reports_warning(tmp_path: Path) -> None:
    path = make_txt(tmp_path, "   \n\n   ", name="empty.txt")
    result = load_document(path)

    assert result.is_empty
    assert result.warning == "未能从文件中提取到文本"


# ---------------- 图片型 PDF ----------------


def test_blank_pdf_is_flagged_needs_ocr(tmp_path: Path) -> None:
    result = load_document(make_blank_pdf(tmp_path, pages=3))

    assert result.needs_ocr is True
    assert "扫描件" in result.warning
    # 空文本的通用告警不应覆盖更具体的 OCR 告警
    assert "未能从文件中提取到文本" not in result.warning


# ---------------- 加密 PDF ----------------


def test_encrypted_pdf_raises_parse_error(tmp_path: Path) -> None:
    with pytest.raises(ParseError, match="已加密"):
        load_document(make_encrypted_pdf(tmp_path))


# ---------------- DOCX ----------------


def test_load_docx_paragraphs_and_tables(tmp_path: Path) -> None:
    result = load_document(make_docx(tmp_path))

    assert not result.is_empty
    text = result.sections[0].text
    assert "第二章 宿舍搬迁" in text
    assert "提前三个工作日提交申请" in text
    # 表格内容不能丢：校园通知大量信息以表格承载
    assert "报到时间" in text
    assert "九月一日至九月二日" in text


# ---------------- 异常分支 ----------------


def test_unsupported_suffix_raises(tmp_path: Path) -> None:
    target = tmp_path / "photo.png"
    target.write_bytes(b"fake")

    with pytest.raises(UnsupportedFileType) as excinfo:
        load_document(target)
    assert ".png" in str(excinfo.value)


def test_uppercase_suffix_is_accepted(tmp_path: Path) -> None:
    target = tmp_path / "NOTICE.TXT"
    target.write_text("校园卡挂失流程。", encoding="utf-8")

    result = load_document(target)
    assert "校园卡挂失流程" in result.sections[0].text


def test_missing_file_raises(tmp_path: Path) -> None:
    with pytest.raises(ParseError, match="不存在"):
        load_document(tmp_path / "not-exist.txt")


def test_directory_path_raises(tmp_path: Path) -> None:
    folder = tmp_path / "docs.txt"
    folder.mkdir()

    with pytest.raises(ParseError, match="不是文件"):
        load_document(folder)


def test_corrupted_pdf_raises_parse_error(tmp_path: Path) -> None:
    target = tmp_path / "broken.pdf"
    target.write_bytes(b"this is definitely not a pdf")

    with pytest.raises(ParseError, match="PDF 解析失败"):
        load_document(target)
